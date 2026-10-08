"""
Generates the Excel fixtures used by the Stage 3 import tests.

 1. tracker-with-history.xlsx  - the real tracker workbook ("Рекомпозиция_трекер.xlsx") filled with the
    athlete's workout history (20.06 - 23.07.2026) and a few nutrition days, recalculated by LibreOffice
    so that formula cells carry cached values, exactly like a workbook saved by Excel.
 2. messy.xlsx                 - a hand-made workbook with another header layout, an unknown sheet and bad rows.

Usage: python3 tools/make-excel-fixtures.py <path to tracker.xlsx>
"""
import datetime as dt
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "packages" / "storage-local" / "test-fixtures"
OUT.mkdir(parents=True, exist_ok=True)

D = dt.datetime
HISTORY = [
    (D(2026, 6, 20), "Ноги", [
        ("Присед в Смитте", [(30, 15)] * 4),
        ("Жим платформы ногами", [(70, 15)] * 4),
        ("Разгибание ног сидя в тренажере", [(50, 15), (50, 15), (55, 15), (55, 15)]),
        ("Сгибание ног лежа в тренажере", [(30, 15), (25, 15), (25, 15), (20, 15)]),
    ]),
    (D(2026, 6, 27), "Грудь, руки, плечи", [
        ("Жим лежа в тренажере Смитта", [(20, 15, "Разминка"), (40, 9), (40, 5), (40, 1.5)]),
        ("Жим гантелей на скамье 30°", [(12, 12), (12, 5), (12, 3)]),
        ("Сведение рук в тренажере", [(25, 13), (25, 10), (25, 10)]),
        ("Подъем руки вверх с нижнего блока (плечи, по очереди)", [(5, 18), (5, 17), (5, 20)]),
        ("Разгибание рук с верхнего блока (по одной руке)", [(10, 20), (15, 7), (15, 6), (15, 4)]),
        ("На предплечья сидя на скамье: раскатывание штанги", [(15, 20), (15, 20), (15, None)]),
    ]),
    (D(2026, 7, 3), "Спина, руки", [
        ("Тяга с верхнего блока узким хватом", [(45.5, 20), (54.5, 10), (63.5, 7), (63.5, 3.5)]),
        ("Сведение в пек-деке", [(27.2, 20), (36.2, 15), (40.8, 7), (40.8, 9)]),
        ("Жим лежа", [(40, 10), (40, 8), (40, 6), (40, 6)]),
        ("Сгибание рук с гантелями", [(12, 10), (12, 6), (10, 7), (10, 10)]),
        ("Французский жим", [(15, 13), (15, 9), (15, 9), (15, 9)]),
    ]),
    (D(2026, 7, 16), "Плечи", [
        ("Подъемы гантелей в стороны", [(12.5, 15), (15, 8), (15, 8), (15, 8)]),
        ("Подъемы гантелей вверх", [(15, 12), (15, 11), (15, 10), (15, 9)]),
        ("Подъем штанги перед собой", [(15, 25)]),
        ("Подъемы вверх сидя в тренажере", [(20, 8), (20, 5), (15, 8), (15, 10)]),
        ("Сидя на задний пучок", [(10, 10), (10, 10), (10, 8), (10, 8)]),
    ]),
    (D(2026, 7, 20), "Руки", [
        ("Сгибание рук с гантелями стоя", [(12.5, 15), (12.5, 15), (12.5, 8), (10, 12)]),
        ("Сгибание рук со штангой стоя", [(25, 8), (25, 8), (25, 12), (25, 8)]),
        ("Разгибание рук из-за головы с гантелей", [(10, 8), (16, 16), (20, 12), (20, 11)]),
        ("Разгибание с верхнего блока (канат)", [(30, 15), (35, 10), (35, 9), (35, 10)]),
        ("Молот", [(10, 9), (10, 5), (9, 6), (9, 6)]),
    ]),
    (D(2026, 7, 23), "Грудь", [
        ("Жим сидя в тренажере", [(30, 15), (35, 8), (35, 8), (35, 8)]),
        ("Жим лежа на наклонной скамье в тренажере Смитта", [(40, 8), (40, 7), (40, 6), (40, 5)]),
        ("Баттерфляй (тренажер)", [(45, 8), (40, 12), (40, 8), (40, 8)]),
    ]),
]

def build_tracker(src: Path) -> Path:
    tmp = OUT / "tracker-with-history.xlsx"
    shutil.copy(src, tmp)
    wb = openpyxl.load_workbook(tmp)
    t = wb["Тренировки"]
    r = 6
    for date, label, exercises in HISTORY:
        for name, sets in exercises:
            for i, s in enumerate(sets, 1):
                weight, reps = s[0], s[1]
                typ = s[2] if len(s) > 2 else "Рабочий"
                t.cell(r, 1, date); t.cell(r, 2, label); t.cell(r, 3, name); t.cell(r, 4, i)
                t.cell(r, 5, weight); t.cell(r, 6, reps); t.cell(r, 8, typ)
                r += 1
    # two deliberately bad rows
    t.cell(r, 1, "31.02.2026"); t.cell(r, 3, "Жим платформы ногами"); t.cell(r, 4, 1); t.cell(r, 5, 80); t.cell(r, 6, 12); r += 1
    t.cell(r, 1, D(2026, 7, 23)); t.cell(r, 3, "Йога"); t.cell(r, 4, 1); t.cell(r, 5, 0); t.cell(r, 6, 30)
    n = wb["Питание"]
    weights = [82.5, 82.3, 82.6, 82.2, 82.1]
    for i, w in enumerate(weights):
        row = 14 + i
        n.cell(row, 2, w); n.cell(row, 3, 2180 - i * 20); n.cell(row, 4, 165); n.cell(row, 5, 70); n.cell(row, 6, 225); n.cell(row, 7, 30); n.cell(row, 8, 2.8)
    n.cell(19, 2, "abc")  # bad weight
    wb.save(tmp)
    subprocess.run([sys.executable, "/mnt/skills/public/xlsx/scripts/recalc.py", str(tmp), "120"], check=False, capture_output=True)
    return tmp

def build_messy() -> Path:
    path = OUT / "messy.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Заметки"
    ws.append(["Купить протеин", "Позвонить в зал"])
    t = wb.create_sheet("Тренировки")
    t.append(["Мой дневник"])
    t.append(["Date", "Exercise", "Set #", "Вес, кг", "Повт.", "RIR", "Тип"])
    t.append([D(2026, 8, 1), "Жим платформы ногами", 1, "80", 12, 2, "Рабочий"])
    t.append(["02.08.2026", "жим платформы ногами", 2, "82,5 кг", "10", "4+", ""])
    t.append([D(2026, 8, 1), "Жим платформы ногами", 1, 80, 12, 2, "Рабочий"])        # duplicate set
    t.append([D(2026, 8, 1), "", 1, 50, 10, None, None])                                # no exercise
    t.append([D(2026, 8, 1), "Румынская тяга", 1, -5, 10, None, None])                  # negative weight
    t.append([D(2026, 8, 1), "Румынская тяга", 2, 20, "много", None, None])             # bad reps
    t.append([D(2026, 8, 1), "Румынская тяга", 3, 20, 10, 9, None])                     # bad rir
    t.append([D(2026, 8, 1), "Румынская тяга", 4, 20, 10, None, "Супер"])               # bad type
    t.append([None, None, None, None, None, None, None])                                # blank row
    t.append([D(2026, 8, 3), "Medball slams", 1, 8, 20, None, None])                    # unknown exercise
    wb.save(path)
    return path

if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/mnt/user-data/outputs/Рекомпозиция_трекер.xlsx")
    print(build_tracker(src)); print(build_messy())
